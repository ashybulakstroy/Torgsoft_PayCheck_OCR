from pathlib import Path
import json

import openpyxl
import pytest

from App.ocr_catalog_eval import (CatalogMatcher, PASS_SIMILARITY, catalog_conflicts, choose_row_boundaries,
                                  classify_name, comparison_key, fallback_variants, load_catalog, normalize_name,
                                  parse_rows, penalty_numbers, punctuation_key, similarity_color, needs_review,
                                  write_excel)
from App.ocr_catalog_eval import aligned_lines


def entry(name, article="1"):
    return dict(name=name, normalized=normalize_name(name), sources=[dict(article=article)])


def test_comparison_key_tolerates_diameter_and_o_after_digit():
    assert comparison_key("Коронка по металлу 025") == comparison_key("Коронка по металлу Ø25")
    assert comparison_key("ПНД Адаптер 4о/32 напа") == comparison_key("ПНД Адаптер 40/32 напа")
    assert comparison_key("ГАЙКА М8 ( 25О ШТ )") == comparison_key("ГАЙКА М8 ( 250 ШТ )")


def test_low_similarity_row_uses_tolerant_comparison():
    matcher = CatalogMatcher([entry("Гайка М8 (250 шт в 1 кг)")])
    result = matcher.match("ГАЙКА М8 ( 25О ШТ В 1 KΓ)")
    assert result["best"]["score"] >= PASS_SIMILARITY
    assert result["status"] == "suggested"


def test_tolerant_comparison_does_not_change_green_rows():
    matcher = CatalogMatcher([entry("Коронка по металлу Ø25")])
    result = matcher.match("Коронка по металлу 025")
    assert result["status"] == "suggested"
    assert result["best"]["score"] < 100


def test_punctuation_key_ignores_missing_decimal_separator():
    assert punctuation_key("Шуруп Пресшайба Сверло 4 2х16") == punctuation_key("Шуруп Пресшайба Сверло 4.2х16")
    assert punctuation_key("Кабель ВВГ нг 3х2 5") == punctuation_key("Кабель ВВГ нг 3х2,5")


def test_low_similarity_ignores_missing_decimal_separator():
    matcher = CatalogMatcher([entry("Шуруп Пресшайба Сверло 4.2х16")])
    result = matcher.match("Шуруп Пресшайба Сверло 4 2х16")
    assert result["status"] == "exact_name"
    assert result["best"]["score"] == 100.0


def test_matching_equates_latin_y_and_cyrillic_u():
    matcher = CatalogMatcher([entry("Клей Универсальный")])
    result = matcher.match("Клей Yниверсальный")
    assert result["status"] == "exact_name"
    assert result["best"]["score"] == 100.0


def test_letter_three_ignores_numeric_penalty_without_equating_names():
    matcher = CatalogMatcher([entry("Штыковая лопата ЗУБР МАСТЕР, ЛКО, без черенка")])
    result = matcher.match("Штыковая лопата 3YБР МАСТEР, ЛКО без черенка")
    assert result["status"] == "suggested"
    assert result["best"]["score"] == 97.67
    assert result["best"]["score"] == result["best"]["raw_similarity"]
    assert result["best"]["missing_query_numbers"] == []


def test_penalty_numbers_keep_real_three_and_skip_codes():
    assert penalty_numbers("3yбp") == set()
    assert penalty_numbers("3yбp 3") == {"3"}
    assert penalty_numbers("3x2.5") == {"3", "2.5"}
    assert penalty_numbers("3 шt") == {"3"}
    assert penalty_numbers("3d 3m") == {"3"}


def test_letter_three_penalty_is_not_applied_to_green_rows():
    matcher = CatalogMatcher([entry("ЛЕПЕСТКОВЫЙ ДИСК КЛТ1 Ø125 (40) ЗУБР")])
    result = matcher.match("ЛЕПЕСТКОВЫЙ ДИСК КЛТ1 125 (40) 3УБP")
    assert result["status"] == "review"
    assert result["best"]["missing_query_numbers"] == ["3"]
    assert result["best"]["score"] < result["best"]["raw_similarity"]


def test_quantity_three_still_gets_numeric_penalty():
    matcher = CatalogMatcher([entry("Шуруп 2 шт")])
    result = matcher.match("Шуруп 3 шт")
    assert result["best"]["missing_query_numbers"] == ["3"]
    assert result["best"]["score"] < result["best"]["raw_similarity"]


def test_low_similarity_prefers_closer_word_count():
    matcher = CatalogMatcher([
        entry('Бита для шуруповерта тип хвостовика C 1/4", № 1, L=25мм, 10шт STAYER PROF', "1"),
        entry("Rodex Отрезной Диск по Металлу 125*1,0", "2"),
    ])
    result = matcher.match("Rоде Отрей днс no MerAnny 128*1.0")
    assert result["best"]["name"] == "Rodex Отрезной Диск по Металлу 125*1,0"


def test_word_count_does_not_override_much_better_match():
    matcher = CatalogMatcher([
        entry("Гровер Ø8 (200шт в 1кг)"),
        entry("2 в 1 Шуруповерт-болгарка  Аккумуляторный"),
    ])
    result = matcher.match("Гровер 08 (20Ошт в 1Kr)")
    assert result["best"]["name"] == "Гровер Ø8 (200шт в 1кг)"
    assert result["status"] == "suggested"


def test_fallback_variants_drop_first_word_and_join_hyphenated():
    variants = fallback_variants("Мойка высокого дав- ления CROWN")
    assert "высокого дав- ления CROWN" in variants
    assert "Мойка высокого давления CROWN" in variants
    assert "высокого давления CROWN" in variants


def test_low_similarity_drops_bad_first_word_and_finds_containment():
    matcher = CatalogMatcher([entry("Алкидная Эмаль ПФ-115 Белый 20 л ELEGANT ( 00158 )")])
    result = matcher.match("n-115Enw20n ELEGANT (00158)")
    assert result["best"]["score"] == 100.0
    assert result["status"] == "suggested"


def test_low_similarity_joins_hyphenated_word():
    matcher = CatalogMatcher([entry("Мойка высокого давления CROWN CT42075")])
    result = matcher.match("Мойка высокого дав- ления CROWN")
    assert result["best"]["score"] == 100.0
    assert result["status"] == "suggested"


def test_matching_normalizes_lookalikes_but_preserves_size_variant():
    matcher = CatalogMatcher([entry("Труба профильная 40х20х2", "2"), entry("Труба профильная 40х20х1.5", "15")])
    result = matcher.match("Труба профильная 40x20*2")
    assert result["status"] == "exact_name"
    assert result["best"]["sources"][0]["article"] == "2"
    assert result["best"]["missing_query_numbers"] == []
    wrong = matcher.match("Труба профильная 40x20x3")
    assert wrong["status"] == "review"
    assert wrong["best"]["missing_query_numbers"]


def test_partial_generic_name_is_not_a_verified_product():
    result = CatalogMatcher([entry("Труба профильная 40х20"), entry("Труба круглая 20")]).match("Труба")
    assert result["status"] == "review"


def test_tax_detection_does_not_exclude_windstopper_product():
    assert classify_name("Костюм Виндстопер осенний") == "product"
    assert classify_name("НДС 16 %") == "tax"
    assert classify_name("HDC 16 %") == "tax"
    assert classify_name("HDС 16 %") == "tax"
    assert classify_name("Тауардын атаув") == "header"
    assert classify_name("Доставка Рынок") == "service"


def test_catalog_reads_names_and_articles_without_using_stock(tmp_path: Path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Название", "Артикул", "Кол-во"])
    sheet.append(["Цемент М-500, 50кг", "000911", 999])
    path = tmp_path / "reference.xlsx"
    workbook.save(path)
    workbook.close()
    catalog, sources = load_catalog(tmp_path)
    assert len(catalog) == 1
    assert catalog[0]["name"] == "Цемент М-500, 50кг"
    assert catalog[0]["sources"][0]["article"] == "000911"
    assert catalog[0]["sources"][0]["row"] == 2
    assert sources[0]["rows"] == 1
    assert "999" not in catalog[0]["normalized"]


def test_separate_zero_in_price_does_not_split_wrapped_product():
    layout = dict(row_boundaries=[16., 97.5, 158.5], name_left=30., name_right=250.,
                  quantity_right=300, right=420., body_top=16, body_bottom=159)
    lines = [dict(text="16 800,0", center=[350, 110]), dict(text="0", center=[350, 135])]
    boundaries, mode = choose_row_boundaries(lines, layout)
    assert boundaries == layout["row_boundaries"]
    assert mode == "grid"


def make_line(text, x, y0, y1, confidence=.9):
    box = [[40., y0], [160., y0], [160., y1], [40., y1]]
    return dict(text=text, center=[x, (y0+y1)/2], box=box, box_original=box, confidence=confidence)


def test_boundary_line_follows_row_anchor_not_text_center():
    layout = dict(name_left=30., name_right=180., quantity_right=220., right=300.)
    lines = [
        make_line("Батарейки DURAcELL", 101., 81., 104.),
        make_line("ORIQINAL (AAA) -2 WT", 102., 92., 116.),
        make_line("680T", 245., 95., 110.),
        make_line("БИТА ДЛЯ ШУРУПО-", 97., 105., 127.),
        make_line("ВЕРТА КОВАННЫЕ 10", 101., 116., 139.),
        make_line("10", 196., 125., 137.),
        make_line("500T", 245., 124., 138.),
        make_line('шТ ЗУБР "ЭКСПЕРТ"', 97., 128., 149.),
    ]
    rows = parse_rows(lines, layout, [80., 116.75, 145.75])
    assert rows[0]["name_raw"] == "Батарейки DURAcELL ORIQINAL (AAA) -2 WT"
    assert rows[1]["name_raw"].startswith("БИТА ДЛЯ ШУРУПО- ВЕРТА КОВАННЫЕ 10")


def test_boundary_line_stays_when_its_anchor_is_closer():
    layout = dict(name_left=30., name_right=180., quantity_right=220., right=300.)
    lines = [
        make_line("ПЕРВЫЙ ТОВАР", 100., 85., 100.),
        make_line("ПОГРАНИЧНАЯ СТРОКА", 100., 101., 113.),
        make_line("100T", 245., 98., 108.),
        make_line("ВТОРОЙ ТОВАР", 100., 115., 135.),
        make_line("200T", 245., 132., 142.),
    ]
    rows = parse_rows(lines, layout, [80., 110., 150.])
    assert "ПОГРАНИЧНАЯ СТРОКА" in rows[0]["name_raw"]
    assert "ПОГРАНИЧНАЯ СТРОКА" not in rows[1]["name_raw"]


@pytest.mark.parametrize("label", ["НДС 16 %", "НДС16%", "HDC", "HDС 16 %"])
def test_tax_cell_keeps_its_numbers_without_splitting_previous_product(label):
    layout = dict(name_left=30., name_right=180., quantity_right=220., right=300.)
    lines = [make_line("Дисковая Пила", 100., 85., 105.),
             make_line("CROWN CT15074", 100., 144., 168.),
             make_line("1", 200., 110., 135.), make_line("31600T", 245., 110., 135.),
             make_line(label, 100., 183., 213.),
             make_line("5056", 200., 181., 209.), make_line("1T", 245., 181., 209.)]
    rows = parse_rows(lines, layout, [80., 158.25, 220.])
    assert [row["kind"] for row in rows] == ["product", "tax"]
    assert rows[0]["name_raw"] == "Дисковая Пила CROWN CT15074"
    assert (rows[0]["quantity_raw"], rows[0]["price_raw"]) == ("1", "31600T")
    assert (rows[1]["quantity_raw"], rows[1]["price_raw"]) == ("5056", "1T")
    assert rows[0]["y_range"][1] == rows[1]["y_range"][0]
    assert rows[0]["y_range"][1] > 168


@pytest.mark.parametrize("variant", ["rectified", "original_crop"])
@pytest.mark.parametrize("stem,expected_suffix,tax_quantity,bad_boundaries,layout", [
    ("WhatsApp Image 2026-10-24 at 13.33.50", "CT15074", "5056", [16., 80., 158.25, 219.],
     dict(name_left=72., name_right=348., quantity_right=436., right=560.)),
    ("WhatsApp Image 2026-10-25 at 13.16.05", "50951", "88656", [16., 89., 220., 360.5, 439.],
     dict(name_left=94., name_right=452., quantity_right=570., right=749.)),
])
def test_real_tax_cell_boundary(stem, expected_suffix, tax_quantity, bad_boundaries, layout, variant):
    directory = Path(__file__).resolve().parents[2] / "eval_outputs" / stem
    payload = json.loads((directory / f"ocr_{variant}.json").read_text(encoding="utf-8"))
    metadata = json.loads((directory / "rectification.json").read_text(encoding="utf-8"))
    lines = aligned_lines(payload, metadata, original=variant == "original_crop")
    rows = parse_rows(lines, layout, bad_boundaries)
    products = [row for row in rows if row["kind"] == "product"]
    tax = [row for row in rows if row["kind"] == "tax"]
    assert len(products) == (1 if expected_suffix == "CT15074" else 2)
    assert products[-1]["name_raw"].endswith(expected_suffix)
    assert len(tax) == 1 and tax[0]["quantity_raw"] == tax_quantity
    assert tax[0]["price_raw"].replace(" ", "").casefold() in {"1t", "1т"}
    assert len(rows) == len(products) + 1


def test_excel_preserves_article_zeros_and_literal_ocr_text(tmp_path):
    row = dict(file="photo.jpeg", row=2, item_number="1", kind="product", ocr_name="=literal OCR",
               quantity="5", price="2900T", reference_name="Цемент", articles="000349",
               similarity=100., status="exact_name", before_similarity=85.)
    path = tmp_path / "matches.xlsx"
    write_excel([row], path)
    workbook = openpyxl.load_workbook(path)
    sheet = workbook.active
    assert sheet["I2"].value == "000349"
    assert sheet["I2"].data_type == "s"
    assert sheet["E2"].value == "=literal OCR"
    assert sheet["E2"].data_type == "s"
    assert sheet["J2"].value == 100
    workbook.close()


def test_similarity_color_threshold_is_inclusive_and_keeps_review_status():
    row = dict(kind="product", similarity=PASS_SIMILARITY, status="review", articles="1")
    assert similarity_color(row) == "catalog-high"
    row["similarity"] = PASS_SIMILARITY - 0.01
    assert similarity_color(row) == "catalog-low"
    row["similarity"] = PASS_SIMILARITY
    assert needs_review(row)  # Color does not approve ambiguous candidates.
    assert row["status"] == "review"
    assert similarity_color(dict(row, kind="tax", similarity=100.)) == ""


def test_excel_confidence_percent_and_color_match_exported_scores(tmp_path):
    base = dict(file="photo.jpeg", row=2, item_number="1", kind="product", ocr_name="Цемент",
                ocr_name_confidence_pct=88.5, quantity="5", price="2900T", reference_name="Цемент",
                articles="000349", similarity=PASS_SIMILARITY - 0.01, status="review", before_similarity=85., review_reason="Проверить")
    rows = [base, dict(base, similarity=PASS_SIMILARITY), dict(base, similarity=100., status="exact_name")]
    path = tmp_path / "report.xlsx"
    write_excel(rows, path)
    workbook = openpyxl.load_workbook(path)
    sheet = workbook["Сопоставление"]
    assert sheet["A3"].fill.fgColor.rgb == sheet["A4"].fill.fgColor.rgb
    assert sheet["A2"].fill.fgColor.rgb != sheet["A3"].fill.fgColor.rgb
    assert sheet["F2"].value == 88.5
    assert sheet["F2"].number_format == '0.00"%"'
    assert workbook["Проверить"].max_row == 3
    workbook.close()


def test_catalog_detects_header_after_metadata_and_article_column(tmp_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = "Состояние склада"
    sheet["A7"] = "Товар"
    sheet["F7"] = "Артикул"
    sheet["A8"] = "Мойка высокого давления CROWN CT42075"
    sheet["F8"] = "9759122"
    path = tmp_path / "priced.xlsx"
    workbook.save(path)
    workbook.close()

    catalog, workbooks = load_catalog(tmp_path)

    assert len(catalog) == 1
    assert catalog[0]["name"] == "Мойка высокого давления CROWN CT42075"
    assert catalog[0]["sources"][0] == dict(workbook="priced.xlsx", sheet="Sheet", row=8, article="9759122")
    assert workbooks[0]["sheets"][0] == dict(sheet="Sheet", rows=1, header_row=7, name_column=0,
                                             article_column=5, mode="header", warning=None)


def test_catalog_falls_back_to_positional_columns_without_header(tmp_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["026 Насос скважинный SN 50/22", "9782267", 1])
    sheet.append(["Мойка высокого давления CROWN CT42075", "9759122", 1])
    path = tmp_path / "third_floor.xlsx"
    workbook.save(path)
    workbook.close()

    catalog, workbooks = load_catalog(tmp_path)

    assert {entry["name"] for entry in catalog} == {
        "026 Насос скважинный SN 50/22",
        "Мойка высокого давления CROWN CT42075",
    }
    assert catalog[0]["sources"][0]["article"] == "9782267"
    assert catalog[0]["sources"][0]["row"] == 1
    assert workbooks[0]["sheets"][0]["mode"] == "positional"
    assert workbooks[0]["sheets"][0]["header_row"] is None


def test_catalog_merges_duplicates_and_reports_only_conflicting_articles(tmp_path):
    same = tmp_path / "same"
    same.mkdir()
    for filename in ("a.xlsx", "b.xlsx"):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Название", "Артикул"])
        sheet.append(["Товар X", "1"])
        workbook.save(same / filename)
        workbook.close()

    catalog, workbooks = load_catalog(same)
    assert len(workbooks) == 2
    assert len(catalog) == 1
    assert len(catalog[0]["sources"]) == 2
    assert catalog_conflicts(catalog) == []

    different = tmp_path / "different"
    different.mkdir()
    for filename, article in (("a.xlsx", "1"), ("b.xlsx", "2")):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Название", "Артикул"])
        sheet.append(["Товар X", article])
        workbook.save(different / filename)
        workbook.close()

    catalog, _ = load_catalog(different)
    conflicts = catalog_conflicts(catalog)
    assert len(conflicts) == 1
    assert conflicts[0]["articles"] == ["1", "2"]


def test_catalog_skips_unusable_sheet_and_keeps_valid_one(tmp_path):
    workbook = openpyxl.Workbook()
    bad = workbook.active
    bad.title = "bad"
    bad.append(["Состояние склада"])
    good = workbook.create_sheet("good")
    good.append(["Название", "Артикул"])
    good.append(["Цемент", "000911"])
    path = tmp_path / "mixed.xlsx"
    workbook.save(path)
    workbook.close()

    catalog, workbooks = load_catalog(tmp_path)

    assert [entry["name"] for entry in catalog] == ["Цемент"]
    assert workbooks[0]["rows"] == 1
    assert any("skipped" in warning for warning in workbooks[0]["warnings"])
    assert [sheet["mode"] for sheet in workbooks[0]["sheets"]] == ["skipped", "header"]


def test_merged_quantity_and_price_split_by_separator():
    layout = dict(name_left=30., name_right=180., quantity_right=220., right=300.)
    lines = [
        make_line("Электрогенератор", 100., 85., 105.),
        dict(text="1| 242000 т", center=[245., 95.],
             eval_box=[[210., 90.], [280., 90.], [280., 110.], [210., 110.]],
             box=[[210., 90.], [280., 90.], [280., 110.], [210., 110.]],
             box_original=[[210., 90.], [280., 90.], [280., 110.], [210., 110.]],
             confidence=0.85)
    ]
    rows = parse_rows(lines, layout, [80., 120.])
    assert len(rows) == 1
    assert rows[0]["name_raw"] == "Электрогенератор"
    assert rows[0]["quantity_raw"] == "1"
    assert rows[0]["price_raw"] == "242000 т"


def test_fallback_split_when_quantity_is_empty():
    layout = dict(name_left=30., name_right=180., quantity_right=220., right=300.)
    lines = [
        make_line("Товар Тестовый", 100., 85., 105.),
        make_line("2| 50000 T", 250., 90., 110.)
    ]
    rows = parse_rows(lines, layout, [80., 120.])
    assert len(rows) == 1
    assert rows[0]["quantity_raw"] == "2"
    assert rows[0]["price_raw"] == "50000 T"


def test_normal_price_with_thousands_separator_is_not_split():
    layout = dict(name_left=30., name_right=180., quantity_right=220., right=300.)
    lines = [
        make_line("Товар Тестовый", 100., 85., 105.),
        make_line("1", 200., 90., 110.),
        make_line("16 800,00 T", 260., 90., 110.)
    ]
    rows = parse_rows(lines, layout, [80., 120.])
    assert len(rows) == 1
    assert rows[0]["quantity_raw"] == "1"
    assert rows[0]["price_raw"] == "16 800,00 T"
