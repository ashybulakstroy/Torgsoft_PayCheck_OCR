"""Compare raw OCR product rows with local XLSX catalogs, without correcting OCR."""
from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import os
import re
import unicodedata
from collections import Counter
from decimal import Decimal
from pathlib import Path

import cv2
import numpy as np
import openpyxl
from dotenv import load_dotenv
from openpyxl.styles import PatternFill
from rapidfuzz import fuzz, process
if __package__:
    from .manual_mappings import ManualMappings
else:
    from manual_mappings import ManualMappings

ROOT = Path(__file__).resolve().parents[1]
load_dotenv(ROOT / ".env")


def env_number(name, default):
    try:
        return float(os.getenv(name, str(default)).strip().rstrip("%"))
    except ValueError:
        return float(default)


PASS_SIMILARITY = env_number("SIMILARITY_PASS_THRESHOLD", 87.0)
STATUS_LABELS = dict(exact_name="Совпало название", suggested="Вероятный вариант", review="Проверить",
                     manual_mapping="Ручное сопоставление из базы знаний",
                     not_found="Кандидат не найден", unreadable="Не прочитано", excluded_tax="НДС",
                     excluded_service="Доставка", excluded_total="Итог", excluded_unreadable="Не прочитано")


def similarity_color(row):
    if row["kind"] != "product":
        return ""
    score = row.get("similarity", "")
    return "catalog-high" if score != "" and float(score) >= PASS_SIMILARITY else "catalog-low"


def needs_review(row):
    if row.get("status") == "manual_mapping":
        return False
    return row["kind"] == "product" and (similarity_color(row) == "catalog-low"
        or row["status"] in {"review", "not_found", "unreadable"}
        or len([a for a in row.get("articles", "").split(";") if a.strip()]) > 1)


def review_reason(row):
    if row["kind"] != "product":
        return ""
    match = row["match"]
    if match.get("source") == "manual_mapping":
        return ""
    best = match["best"]
    if best is None:
        return "Название не прочитано или кандидат не найден"
    reasons = []
    if best["score"] < PASS_SIMILARITY:
        reasons.append(f"Сходство ниже {PASS_SIMILARITY:g}%")
    if best.get("missing_query_numbers"):
        missing = ", ".join(format(Decimal(value), "f") for value in best["missing_query_numbers"])
        reasons.append("Не совпали числа в названии: " + missing)
    if match["status"] != "exact_name" and match.get("gap", 0) < 5:
        reasons.append("Близкие варианты в справочнике")
    if len({source["article"] for source in best["sources"] if source["article"]}) > 1:
        reasons.append("У названия несколько артикулов")
    if match["status"] == "review" and not reasons:
        reasons.append("Название требует проверки")
    return "; ".join(reasons)


def normalize_name(value):
    text = unicodedata.normalize("NFKC", str(value)).casefold().replace("ё", "е")
    text = re.sub(r"(?<=\d)[,](?=\d)", ".", text)
    text = re.sub(r"(?<=\d)\s*[xх×*]\s*(?=\d)", "x", text)
    text = re.sub(r"[^\w\d.ø⌀]+", " ", text, flags=re.UNICODE)
    # A matching key only: raw OCR is always preserved separately. Latin and
    # Cyrillic lookalikes are common in mixed brand names and printed capitals.
    text = text.translate(str.maketrans("авсенкмортху", "abcehkmoptxy"))
    return " ".join(text.split())


def comparison_key(value):
    """Tolerant key for scoring low-similarity rows; stored names stay unchanged."""
    text = normalize_name(value)
    text = re.sub(r"[ø⌀∅]", "", text)
    text = re.sub(r"(?<=\d)[oо]", "0", text)
    text = re.sub(r"(?<!\d)0+(?=\d)", "", text)
    return " ".join(text.split())


def punctuation_key(value):
    """Comparison-only key where dots/commas are optional in OCR text."""
    text = comparison_key(value)
    text = text.replace(".", "").replace(",", "")
    merged = []
    for token in text.split():
        if merged and re.fullmatch(r"\d+", merged[-1]) and re.fullmatch(r"\d+x\d+", token):
            merged[-1] += token
        elif merged and re.fullmatch(r"\d+x\d+", merged[-1]) and re.fullmatch(r"\d+", token):
            merged[-1] += token
        else:
            merged.append(token)
    return " ".join(merged)


def word_count(text):
    """Count word-like tokens, ignoring standalone numbers and codes."""
    return sum(1 for token in text.split() if re.search(r"[^\W\d_]", token, re.UNICODE))


def word_count_factor(query, candidate):
    """Rank viable variants by how close their word counts are."""
    query_words = word_count(query)
    candidate_words = word_count(candidate)
    if not query_words or not candidate_words:
        return 1.0
    difference = abs(query_words - candidate_words) / max(query_words, candidate_words)
    return 1.0 - 0.35 * difference


HYPHEN_BREAK = re.compile(r"(?<=\w)-\s+(?=\w)")


def fallback_variants(value):
    """Raw OCR variants for a low-similarity row; original text is not changed."""
    text = " ".join(str(value).split())
    words = text.split()
    variants = []
    if len(words) >= 3:
        variants.append(" ".join(words[1:]))
    joined = HYPHEN_BREAK.sub("", text)
    if joined != text:
        variants.append(joined)
        joined_words = joined.split()
        if len(joined_words) >= 3:
            variants.append(" ".join(joined_words[1:]))
    seen, result = set(), []
    for variant in variants:
        key = comparison_key(variant)
        if key not in seen:
            seen.add(key)
            result.append(variant)
    return result


def numbers(text):
    return {str(Decimal(n.replace(",", ".")).normalize()) for n in re.findall(r"\d+(?:[.,]\d+)?", text)}


def penalty_numbers(value):
    """Numbers for scoring; a leading 3 in a word token is an OCR lookalike for З."""
    tokens = []
    for token in value.split():
        match = re.fullmatch(r"3([^\W\d_]{3,})", token)
        if match and re.search(r"[а-я]", match.group(1)):
            token = "з" + token[1:]
        tokens.append(token)
    return numbers(" ".join(tokens))


NAME_HEADERS = {"название", "наименование", "наименование товара", "товар", "name", "product"}
ARTICLE_HEADERS = {"артикул", "артикул товара", "арт", "код", "sku", "article"}
SKIP_CATALOG_NAMES = {"итого", "всего", "товар", "название", "наименование", "артикул", "состояние склада"}
HEADER_SCAN_ROWS = 50


def _cell_text(value):
    return "" if value is None else str(value).strip()


def _find_header(row):
    cells = [_cell_text(value).casefold() for value in row]
    name_column = next((i for i, value in enumerate(cells) if value in NAME_HEADERS), None)
    article_column = next((i for i, value in enumerate(cells) if value in ARTICLE_HEADERS), None)
    return name_column, article_column


def _looks_like_positional_row(row):
    if not row or row[0] is None:
        return False
    name = _cell_text(row[0])
    if normalize_name(name) in SKIP_CATALOG_NAMES:
        return False
    filled = sum(1 for value in row if value is not None and _cell_text(value))
    return filled >= 2 and any(ch.isalpha() for ch in name)


def catalog_conflicts(entries):
    conflicts = []
    for entry in entries:
        articles = sorted({source["article"] for source in entry["sources"] if source["article"]})
        if len(articles) > 1:
            conflicts.append(dict(name=entry["name"], normalized=entry["normalized"], articles=articles,
                                  sources=entry["sources"]))
    return conflicts


def load_catalog(directory):
    grouped, workbooks = {}, []
    for path in sorted(Path(directory).glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        try:
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            workbooks.append(dict(file=path.name, rows=0, sha256="", sheets=[], warnings=[f"failed to open: {exc}"]))
            continue
        count = 0
        sheet_meta, warnings = [], []
        try:
            for sheet in workbook:
                rows = sheet.iter_rows(values_only=True)
                buffered = []
                name_column = article_column = None
                header_row = None
                mode = "header"
                for row_number, row in enumerate(rows, 1):
                    if row_number <= HEADER_SCAN_ROWS:
                        candidate_name, candidate_article = _find_header(row)
                        if candidate_name is not None:
                            name_column = candidate_name
                            article_column = candidate_article
                            header_row = row_number
                            break
                    buffered.append((row_number, row))
                    if row_number >= HEADER_SCAN_ROWS:
                        break
                if name_column is None:
                    if any(_looks_like_positional_row(row) for _, row in buffered):
                        name_column, article_column = 0, 1
                        mode = "positional"
                        warning = f"{path.name}/{sheet.title}: header not found; positional A/B"
                        warnings.append(warning)
                    else:
                        warning = f"{path.name}/{sheet.title}: skipped (no usable header or data)"
                        warnings.append(warning)
                        sheet_meta.append(dict(sheet=sheet.title, rows=0, header_row=None, name_column=None,
                                               article_column=None, mode="skipped", warning=warning))
                        continue
                else:
                    warning = None
                sheet_count = 0

                def add_row(row_number, row):
                    nonlocal sheet_count
                    if not row or name_column >= len(row) or row[name_column] is None:
                        return
                    name = str(row[name_column]).strip()
                    normalized = normalize_name(name)
                    if len(normalized) < 2 or not any(c.isalpha() for c in normalized):
                        return
                    if normalized in SKIP_CATALOG_NAMES:
                        return
                    article = ""
                    if article_column is not None and article_column < len(row) and row[article_column] is not None:
                        article = str(row[article_column]).strip()
                    entry = grouped.setdefault(normalized, dict(name=name, normalized=normalized, sources=[]))
                    entry["sources"].append(dict(workbook=path.name, sheet=sheet.title, row=row_number, article=article))
                    sheet_count += 1

                if mode == "positional":
                    for row_number, row in buffered:
                        add_row(row_number, row)
                    start_row = len(buffered) + 1
                else:
                    start_row = header_row + 1
                for row_number, row in enumerate(rows, start_row):
                    add_row(row_number, row)
                count += sheet_count
                sheet_meta.append(dict(sheet=sheet.title, rows=sheet_count, header_row=header_row,
                                       name_column=name_column, article_column=article_column,
                                       mode=mode, warning=warning))
        finally:
            workbook.close()
        workbooks.append(dict(file=path.name, rows=count, sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                              sheets=sheet_meta, warnings=warnings))
    if not grouped:
        raise ValueError("Reference catalog is empty; comparison is impossible")
    return list(grouped.values()), workbooks


class CatalogMatcher:
    def __init__(self, entries, knowledge=None):
        self.entries = entries
        self.names = [entry["normalized"] for entry in entries]
        self.comparison_names = [comparison_key(entry["normalized"]) for entry in entries]
        self.punctuation_names = [punctuation_key(entry["normalized"]) for entry in entries]
        self.cache = {}
        self.knowledge = knowledge if knowledge is not None else ManualMappings()
        self.catalog_index = {entry["normalized"]:entry for entry in entries}

    def _match(self, query, names, rank_by_words=False, letter_three=False):
        number_fn = penalty_numbers if letter_three else numbers
        query_numbers = number_fn(query)
        extracted = list(process.extract(query, names, scorer=fuzz.WRatio,
                                         limit=50 if rank_by_words else 15))
        if rank_by_words:
            query_words = word_count(query)
            closest = min(abs(query_words - word_count(names[index])) for _, _, index in extracted)
            if closest:
                extracted = extracted[:15]
                rank_by_words = False
        candidates = []
        for _, raw_score, index in extracted:
            entry = self.entries[index]
            entry_name = names[index]
            missing = query_numbers - number_fn(entry_name)
            penalty = 1 - .15 * len(missing) / max(1, len(query_numbers))
            rank_penalty = word_count_factor(query, entry_name) if rank_by_words and query != entry_name else 1.0
            candidates.append(dict(name=entry["name"], score=round(raw_score*penalty*rank_penalty, 2),
                                   raw_similarity=round(raw_score, 2), missing_query_numbers=sorted(missing),
                                   normalized=entry["normalized"], sources=entry["sources"], _name=entry_name))
        candidates.sort(key=lambda c: c["score"], reverse=True)
        best = candidates[0]
        gap = best["score"] - candidates[1]["score"] if len(candidates) > 1 else best["score"]
        if query == best["_name"]:
            status = "exact_name"
        elif best["score"] >= 90 and gap >= 5 and not best["missing_query_numbers"]:
            status = "suggested"
        elif best["score"] >= 55:
            status = "review"
        else:
            status = "not_found"
        for candidate in candidates:
            candidate.pop("_name", None)
        return dict(status=status, best=best, gap=round(gap, 2), top_matches=candidates[:5])

    def _comparison_match(self, query):
        result = self._match(comparison_key(query), self.comparison_names, rank_by_words=True, letter_three=True)
        punctuation = self._match(punctuation_key(query), self.punctuation_names, rank_by_words=True, letter_three=True)
        if punctuation["best"]["score"] > result["best"]["score"]:
            result = punctuation
        return result

    def _word_gap(self, query, result):
        return abs(word_count(query) - word_count(result["best"]["normalized"]))

    def _containment_result(self, query, other, names):
        if len(query) < 10 or len(query.split()) < 2:
            return None
        query_numbers = penalty_numbers(query)
        hits = [index for index, name in enumerate(names)
                if query in name and not query_numbers - penalty_numbers(name)]
        if len(hits) != 1:
            return None
        index = hits[0]
        entry = self.entries[index]
        candidate = dict(name=entry["name"], score=100.0, raw_similarity=100.0, missing_query_numbers=[],
                         normalized=entry["normalized"], sources=entry["sources"])
        others = [c for c in other["top_matches"] if c["name"] != entry["name"]]
        gap = 100.0 - others[0]["score"] if others else 100.0
        if query == names[index]:
            status = "exact_name"
        elif gap >= 5:
            status = "suggested"
        else:
            status = "review"
        return dict(status=status, best=candidate, gap=round(gap, 2), top_matches=[candidate]+others[:4])

    def match(self, raw_text):
        automatic = self._automatic_match(raw_text)
        return self.knowledge.resolve(raw_text, self.catalog_index, automatic) if self.knowledge else automatic

    def _automatic_match(self, raw_text):
        query = normalize_name(raw_text)
        if query in self.cache:
            return self.cache[query]
        if len(query) < 3 or not any(c.isalpha() for c in query):
            return dict(status="unreadable", best=None, top_matches=[])
        result = self._match(query, self.names)
        tolerant_query = comparison_key(raw_text)
        if result["best"]["score"] < PASS_SIMILARITY:
            result = self._match(query, self.names, rank_by_words=True, letter_three=True)
            tolerant = self._comparison_match(tolerant_query)
            if tolerant["best"]["score"] > result["best"]["score"]:
                result = tolerant
        if result["best"]["score"] < PASS_SIMILARITY:
            seen = {query, tolerant_query}
            for variant in fallback_variants(raw_text):
                variant_query = comparison_key(variant)
                if variant_query in seen:
                    continue
                seen.add(variant_query)
                variant_result = self._comparison_match(variant_query)
                if (variant_result["best"]["score"] > result["best"]["score"]
                        and self._word_gap(query, variant_result) <= self._word_gap(query, result)):
                    result = variant_result
                for containment_query, names in ((variant_query, self.comparison_names),
                                                  (punctuation_key(variant_query), self.punctuation_names)):
                    containment = self._containment_result(containment_query, variant_result, names)
                    if containment and containment["best"]["score"] > result["best"]["score"]:
                        result = containment
        self.cache[query] = result
        return result


def table_layout(image, metadata):
    if metadata.get("method")=="wave_grid" and metadata.get("grid_layout"):
        return dict(metadata["grid_layout"],row_boundaries=list(metadata["grid_layout"]["row_boundaries"]))
    margin = metadata["margin"]
    width, height = metadata["body_size"]
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    binary = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, cv2.THRESH_BINARY_INV, 31, 9)
    body = binary[margin:margin+height, margin:margin+width]
    vertical = cv2.morphologyEx(body, cv2.MORPH_CLOSE, np.ones((5, 1), np.uint8))
    vertical = cv2.morphologyEx(vertical, cv2.MORPH_OPEN, np.ones((max(15, int(height*.18)), 1), np.uint8))
    projection = np.count_nonzero(vertical, axis=0)
    def divider(lo, hi, fallback):
        start, end = int(lo*width), int(hi*width)
        index = start + int(np.argmax(projection[start:end]))
        return margin + (index if projection[index] > height*.18 else width*fallback)
    name_left = divider(.035, .18, .10)
    name_right = divider(.45, .74, .61)
    quantity_right = divider(max(.69, (name_right-margin)/width+.06), .91, .80)
    # A narrow name-column probe tolerates small residual line curvature.
    lo, hi = int(name_left+4), int(name_right-4)
    probe = binary[:, lo:hi]
    thick = cv2.dilate(probe, np.ones((3, 1), np.uint8))
    rules = cv2.morphologyEx(thick, cv2.MORPH_OPEN, np.ones((1, max(20, int((hi-lo)*.60))), np.uint8))
    profile = np.count_nonzero(rules, axis=1)
    # The empty space around row numbers often preserves rules lost between
    # long wrapped product names. Do not bridge adjacent letters horizontally.
    number_probe = binary[:, margin+3:max(margin+4, int(name_left)-3)]
    number_rules = cv2.morphologyEx(cv2.dilate(number_probe, np.ones((3, 1), np.uint8)),
                                    cv2.MORPH_OPEN, np.ones((1, max(8, int(number_probe.shape[1]*.8))), np.uint8))
    number_profile = np.count_nonzero(number_rules, axis=1)
    selected = (profile > (hi-lo)*.55) | (number_profile > number_probe.shape[1]*.85)
    positions = np.flatnonzero(selected & (np.arange(len(profile)) >= margin-4)
                               & (np.arange(len(profile)) <= margin+height+4))
    groups = []
    for value in positions:
        if not groups or value-groups[-1][-1] > max(3, min(8, height*.008)):
            groups.append([int(value)])
        else:
            groups[-1].append(int(value))
    boundaries = [float(np.mean(group)) for group in groups]
    if not boundaries or boundaries[0] > margin+8:
        boundaries.insert(0, float(margin))
    if boundaries[-1] < margin+height-8:
        boundaries.append(float(margin+height-1))
    return dict(name_left=float(name_left), name_right=float(name_right), quantity_right=float(quantity_right),
                right=float(margin+width), body_top=margin, body_bottom=margin+height, row_boundaries=boundaries)


def classify_name(text):
    # Classification uses real letters; catalog matching uses a visual key.
    normalized = unicodedata.normalize("NFKC", str(text)).casefold()
    compact = normalized.replace(" ", "")
    if not compact:
        return "unreadable"
    if any(token in normalized for token in ("атау", "атайы", "наименование", "название")):
        return "header"
    if re.match(r"^(?:[нh][дd][сc]|налог)(?:\d|%|$)", compact):
        return "tax"
    if "доставка" in normalized:
        return "service"
    if "корыт" in normalized or "қорыт" in normalized or "итого" in normalized:
        return "total"
    return "product"def is_table_header_row(row):    """Recognize a printed table header after its cells were grouped."""    def compact(value):        return unicodedata.normalize("NFKC", str(value or "")).casefold().replace(" ", "")    name = compact(row.get("name_raw", ""))    quantity = compact(row.get("quantity_raw", ""))    price = compact(row.get("price_raw", ""))    item = compact(row.get("item_number_raw", ""))    markers = [        name in {"n", "no", "ne", "№"} or item in {"n", "no", "ne", "№"}        or any(token in name for token in ("атау", "атайы", "наимен", "назван")),        any(token in quantity for token in ("саны", "колво", "кол-во", "колич")),        any(token in price for token in ("бага", "цена", "стоим")),    ]    return sum(markers) >= 2


def aligned_lines(payload, metadata, original=False):
    lines = []
    if original:
        width, height = metadata["body_size"]
        margin = metadata["margin"]
        target = np.array([[margin, margin], [margin+width-1, margin],
                           [margin+width-1, margin+height-1], [margin, margin+height-1]], np.float32)
        matrix = cv2.getPerspectiveTransform(np.asarray(metadata["source_body_corners"], np.float32), target)
    for line in payload["lines"]:
        box = np.asarray(line["box_original"] if original else line["box"], np.float32)
        if original:
            box = cv2.perspectiveTransform(box[None], matrix)[0]
        lines.append(dict(line, eval_box=box, center=box.mean(axis=0)))
    return lines


def line_y_bounds(line):
    box = line.get("eval_box")
    if box is None:
        box = line.get("box")
    if box is not None:
        values = np.asarray(box, dtype=float)[:, 1]
        return float(values.min()), float(values.max())
    y = float(line["center"][1])
    return y, y


def is_name_line(line, layout):
    return layout["name_left"]-3 <= line["center"][0] < layout["name_right"]


def row_anchor_y(members, layout):
    anchors = [line["center"][1] for line in members
               if layout["name_right"] <= line["center"][0] < layout["right"]+4
               and re.search(r"\d", line["text"])]
    return float(np.median(anchors)) if anchors else None


def resolve_boundary_lines(bands, layout):
    for index, band in enumerate(bands):
        name_lines = [line for line in band["members"]
                      if is_name_line(line, layout)
                      and classify_name(line["text"]) != "header"]
        if len(name_lines) <= 1:
            continue
        for line in list(band["members"]):
            if not is_name_line(line, layout) or classify_name(line["text"]) == "header":
                continue
            top, bottom = line_y_bounds(line)
            center = float(line["center"][1])
            for neighbor_index in (index-1, index+1):
                if not 0 <= neighbor_index < len(bands):
                    continue
                neighbor = bands[neighbor_index]
                boundary = band["start"] if neighbor_index < index else band["end"]
                if not top < boundary < bottom:
                    continue
                if abs(center-boundary) > max(3.0, (bottom-top)*.2):
                    continue
                if band["anchor"] is None or neighbor["anchor"] is None:
                    continue
                if neighbor_index > index and max(name_lines, key=lambda l: l["center"][1]) is not line:
                    continue
                if neighbor_index < index and min(name_lines, key=lambda l: l["center"][1]) is not line:
                    continue
                if abs(bottom-band["anchor"]) <= abs(bottom-neighbor["anchor"])+2:
                    continue
                band["members"].remove(line)
                neighbor["members"].append(line)
                neighbor["members"].sort(key=lambda l: (l["center"][1], l["center"][0]))
                break


def enforce_cell_boundaries(lines, layout, row_boundaries):
    """A tax/service/total label starts a cell, including its numeric columns.

    Price midpoints are only estimates: they can cut the last wrapped line of
    the preceding product. Prefer the whitespace before the observed label.
    """
    boundaries = sorted(set(float(value) for value in row_boundaries))
    if len(boundaries) < 2:
        return boundaries
    names = [line for line in lines if is_name_line(line, layout)
             and boundaries[0] <= line["center"][1] < boundaries[-1]
             and classify_name(line["text"]) != "header"]
    markers = sorted([line for line in names if classify_name(line["text"]) in {"tax", "service", "total"}],
                     key=lambda line: line["center"][1])
    for marker in markers:
        marker_center = float(marker["center"][1])
        preceding = [line for line in names if line["center"][1] < marker_center
                     and line_y_bounds(line)[0] < line_y_bounds(marker)[0]]
        if not preceding:
            continue
        previous = max(preceding, key=lambda line: line["center"][1])
        previous_center = float(previous["center"][1])
        previous_bottom = line_y_bounds(previous)[1]
        marker_top = line_y_bounds(marker)[0]
        if previous_bottom < marker_top:
            safe = [value for value in boundaries[1:-1] if previous_bottom < value < marker_top]
            separator = min(safe, key=lambda value: abs(value-(previous_bottom+marker_top)/2)) if safe else (previous_bottom+marker_top)/2
        else:
            separator = (previous_center+marker_center)/2
        boundaries = [value for index, value in enumerate(boundaries)
                      if index in (0, len(boundaries)-1) or not previous_center < value < marker_center]
        boundaries.append(separator)
        boundaries = sorted(set(boundaries))
    return boundaries


def split_quantity_price_line(line, layout):
    cx = float(line["center"][0])
    name_right = float(layout.get("name_right", 0))
    if cx < name_right - 5:
        return None
    text = str(line.get("text", "")).strip()
    if not text:
        return None

    # Check separator: pipe, exclamation, slash, backslash between digits
    if re.search(r"\d\s*[|!/\\]+\s*\d", text):
        parts = re.split(r"\s*[|!/\\]+\s*", text, maxsplit=1)
        if len(parts) == 2 and re.search(r"\d", parts[0]) and re.search(r"\d", parts[1]):
            return parts[0].strip(), parts[1].strip()

    # Check OCR-confused vertical separator 'I' or 'l' with spaces between digits
    if re.search(r"\d\s+[Il]\s+\d", text):
        parts = re.split(r"\s+[Il]\s+", text, maxsplit=1)
        if len(parts) == 2 and re.search(r"\d", parts[0]) and re.search(r"\d", parts[1]):
            return parts[0].strip(), parts[1].strip()

    return None


def parse_rows(lines, layout, row_boundaries):
    row_boundaries = enforce_cell_boundaries(lines, layout, row_boundaries)
    bands = []
    for index, (start, end) in enumerate(zip(row_boundaries, row_boundaries[1:])):
        if end-start < 8:
            continue
        members = sorted([line for line in lines if start <= line["center"][1] < end], key=lambda l: (l["center"][1], l["center"][0]))
        if not members:
            continue
        bands.append(dict(index=index, start=start, end=end, members=members, anchor=row_anchor_y(members, layout)))
    resolve_boundary_lines(bands, layout)
    rows = []
    for band in bands:
        members = band["members"]
        name, quantities, prices, item_numbers = [], [], [], []
        header_found = False
        for line in members:
            cx = line["center"][0]
            if is_name_line(line, layout):
                if classify_name(line["text"]) == "header":
                    header_found = True
                else:
                    name.append(line)
            else:
                split = split_quantity_price_line(line, layout)
                if split:
                    qty, prc = split
                    quantities.append(qty)
                    prices.append(prc)
                elif layout["name_right"] <= cx < layout["quantity_right"]:
                    quantities.append(line["text"])
                elif layout["quantity_right"] <= cx < layout["right"]+4:
                    prices.append(line["text"])
                elif cx < layout["name_left"]-3:
                    item_numbers.append(line["text"])

        # Fallback if a merged line was categorized solely into price or quantity
        if not quantities and prices:
            new_prices = []
            for p in prices:
                if not quantities and re.search(r"\d\s*[|!/\\]+\s*\d", p):
                    parts = re.split(r"\s*[|!/\\]+\s*", p.strip(), maxsplit=1)
                    if len(parts) == 2 and re.search(r"\d", parts[0]) and re.search(r"\d", parts[1]):
                        quantities.append(parts[0].strip())
                        new_prices.append(parts[1].strip())
                        continue
                new_prices.append(p)
            prices = new_prices

        name_text = " ".join(line["text"] for line in name)
        row = dict(row=len(rows)+1, item_number_raw=" ".join(item_numbers), y_range=[band["start"], band["end"]], name_raw=name_text, quantity_raw=" ".join(quantities),
                   price_raw=" ".join(prices), kind="header" if header_found and not name else classify_name(name_text),
                   ocr_confidence=float(np.mean([l["confidence"] for l in name])) if name else 0,
                   raw_lines=[{k: l[k] for k in ("text", "confidence", "box", "box_original")} for l in members])
        if row["kind"] != "header" and not is_table_header_row(row):
            rows.append(row)
    return rows


def choose_row_boundaries(lines, layout):
    rules = layout["row_boundaries"]
    if layout.get("grid_confirmed"):
        return enforce_cell_boundaries(lines,layout,rules),"verified_printed_grid"
    prices = sorted([line for line in lines if line["center"][0] >= layout["quantity_right"]
                     and layout["body_top"]+12 < line["center"][1] < layout["body_bottom"]
                     and re.search(r"\d", line["text"]) and len(line["text"].strip()) >= 2], key=lambda line: line["center"][1])
    # Grid is preferred. Missing rules can merge two item prices into one band.
    boundaries = list(rules)
    inserted = False
    header_lines = [line for line in lines if classify_name(line["text"]) == "header"
                    and line["center"][1] < layout["body_top"]+.45*(layout["body_bottom"]-layout["body_top"])]
    if header_lines:
        bottom = max(float(np.max(line["eval_box"][:, 1])) for line in header_lines)
        candidate = bottom+4
        if all(abs(candidate-existing) > 10 for existing in boundaries):
            boundaries.append(candidate)
    for start, end in zip(rules, rules[1:]):
        anchors = [float(line["center"][1]) for line in prices if start < line["center"][1] < end]
        for a, b in zip(anchors, anchors[1:]):
            if b-a > 12:
                boundaries.append((a+b)/2)
                inserted = True
    return enforce_cell_boundaries(lines, layout, boundaries), "grid_with_price_splits" if inserted else "grid"


def stats(rows):
    products = [row for row in rows if row["kind"] == "product"]
    scores = [row["match"]["best"]["score"] for row in products if row["match"]["best"]]
    return dict(product_rows=len(products), statuses=dict(Counter(row["match"]["status"] for row in products)),
                mean_similarity=round(float(np.mean(scores)), 2) if scores else 0.)


def write_excel(rows, path):
    workbook = openpyxl.Workbook()
    headers = list(rows[0])
    headings = dict(file="Фото", row="Строка области", item_number="№ из OCR", kind="Тип строки", ocr_name="Название OCR",
                    ocr_name_confidence_pct="Уверенность OCR имени, %", review_reason="Причина проверки",
                    quantity="Количество OCR", price="Цена OCR", reference_name="Название справочника", articles="Артикулы",
                    similarity="Сходство / 100", status="Результат", before_similarity="Сходство до / 100")
    for index, (title, selected) in enumerate((("Сопоставление", rows), ("Проверить", [row for row in rows if needs_review(row)]))):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = title
        sheet.append([headings[key] for key in headers])
        for flat in selected:
            values = dict(flat, status=STATUS_LABELS.get(flat["status"], flat["status"]),
                          kind=dict(product="Товар", tax="НДС", service="Доставка", unreadable="Не прочитано", total="Итог").get(flat["kind"], flat["kind"]))
            sheet.append([values[key] for key in headers])
            color = similarity_color(flat)
            fill = PatternFill("solid", fgColor="E4F5E4" if color == "catalog-high" else "FFF2D3") if color else None
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                    cell.number_format = "@"
                if headers[cell.column-1] == "ocr_name_confidence_pct":
                    cell.number_format = '0.00"%"'
                if fill:
                    cell.fill = fill
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column_index, key in enumerate(headers, 1):
            letter = openpyxl.utils.get_column_letter(column_index)
            sheet.column_dimensions[letter].width = 55 if key in {"file", "ocr_name", "reference_name", "review_reason"} else 23
    workbook.save(path)
    workbook.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--reuse-matches", action="store_true", help="Rebuild reports from saved OCR matches without recognition or matching")
    args = parser.parse_args()
    entries, workbooks = load_catalog(ROOT / "data/reference/torgsoft_ref")
    conflicts = catalog_conflicts(entries)
    print(f"catalog: {len(entries)} names, {len(workbooks)} files, {len(conflicts)} article conflicts", flush=True)
    matcher = CatalogMatcher(entries)
    metadata_rows = json.loads((ROOT / "eval_outputs/rectification_summary.json").read_text(encoding="utf-8"))
    photos, csv_rows, sections = [], [], []
    for metadata in metadata_rows:
        directory = ROOT / "eval_outputs" / Path(metadata["file"]).stem
        if args.reuse_matches:
            report = json.loads((directory / "ocr_reference_matches.json").read_text(encoding="utf-8"))
            before, after = report["before"], report["after"]
            layout, boundaries, segmentation = report["layout"], report["row_boundaries"], report["segmentation"]
        else:
            rectified = json.loads((directory / "ocr_rectified.json").read_text(encoding="utf-8"))
            original = json.loads((directory / "ocr_original_crop.json").read_text(encoding="utf-8"))
            layout = table_layout(cv2.imread(str(directory / "rectified.png")), metadata)
            rectified_lines = aligned_lines(rectified, metadata)
            boundaries, segmentation = choose_row_boundaries(rectified_lines, layout)
            after = parse_rows(rectified_lines, layout, boundaries)
            before = parse_rows(aligned_lines(original, metadata, original=True), layout, boundaries)
        for rows in (before, after):
            for row in rows:
                if not args.reuse_matches:
                    row["match"] = matcher.match(row["name_raw"]) if row["kind"] == "product" else dict(status="excluded_"+row["kind"], best=None, top_matches=[])
                row["ocr_name_confidence_pct"] = round(float(row["ocr_confidence"])*100, 2) if row["name_raw"].strip() else None
        report = dict(file=metadata["file"], segmentation=segmentation, layout=layout, row_boundaries=boundaries,
                      before=before, after=after, original_alignment="approximate_homography",
                      before_stats=stats(before), after_stats=stats(after))
        (directory / "ocr_reference_matches.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        photos.append(report)
        table_rows = []
        for row in after:
            best = row["match"]["best"] or {}
            articles = sorted({s["article"] for s in best.get("sources", []) if s["article"]})
            prior = next((r for r in before if r["row"] == row["row"]), None)
            prior_best = prior["match"]["best"] if prior else None
            flat = dict(file=metadata["file"], row=row["row"], item_number=row["item_number_raw"], kind=row["kind"], ocr_name=row["name_raw"],
                        ocr_name_confidence_pct=row["ocr_name_confidence_pct"] if row["ocr_name_confidence_pct"] is not None else "",
                        quantity=row["quantity_raw"], price=row["price_raw"], reference_name=best.get("name", ""),
                        articles="; ".join(articles), similarity=best.get("score", ""), status=row["match"]["status"],
                        before_similarity=prior_best["score"] if prior_best else "", review_reason=review_reason(row))
            csv_rows.append(flat)
            displayed = dict(flat, status=STATUS_LABELS.get(flat["status"], flat["status"]))
            table_rows.append(f'<tr class="{similarity_color(flat)}" data-kind="{flat["kind"]}" '
                              f'data-pending="{str(needs_review(flat)).lower()}" data-status="{html.escape(flat["status"])}">'+''.join(
                '<td>'+html.escape(str(displayed[key]) if displayed[key] != "" else "—")+'</td>' for key in
                ("item_number", "ocr_name", "ocr_name_confidence_pct", "quantity", "price", "reference_name", "articles", "similarity", "status", "before_similarity", "review_reason"))+'</tr>')
        with (directory / "ocr_reference_matches.csv").open("w", encoding="utf-8-sig", newline="") as stream:
            selected = [r for r in csv_rows if r["file"] == metadata["file"]]
            writer = csv.DictWriter(stream, fieldnames=list(csv_rows[0]))
            writer.writeheader(); writer.writerows(selected)
        stem = Path(metadata["file"]).stem
        stat = report["after_stats"]
        statuses = stat["statuses"]
        caption = (f'Товарных строк: {stat["product_rows"]}. Совпало названий: {statuses.get("exact_name", 0)}; '
                   f'вероятных вариантов: {statuses.get("suggested", 0)}; проверить: {statuses.get("review", 0)}. '
                   f'Среднее сходство до / после: {report["before_stats"]["mean_similarity"]} / {stat["mean_similarity"]}.')
        sections.append(f'<details class="photo" open><summary>{html.escape(stem)}</summary><p>{html.escape(caption)}</p>'
                        f'<a href="{html.escape(stem)}/rectified.png">Выпрямленная таблица</a> · '
                        f'<a href="{html.escape(stem)}/ocr_rectified.txt">Исходный текст OCR</a>'
                        '<div class="scroll"><table><thead><tr><th>№ из OCR</th><th>OCR без исправлений</th>'
                        '<th title="Среднее значение confidence OCR-блоков, относящихся только к названию товара, ×100">Уверенность OCR имени, %</th>'
                        '<th>Кол-во OCR</th><th>Цена OCR</th><th>Кандидат справочника</th><th>Артикулы</th><th>Сходство /100</th><th>Статус</th><th>До выпрямления /100</th><th>Причина проверки</th></tr></thead><tbody>'
                        + ''.join(table_rows)+'</tbody></table></div></details>')
        print(f'{metadata["file"]}: {report["after_stats"]}', flush=True)
    totals = Counter()
    for photo in photos:
        totals.update(photo["after_stats"]["statuses"])
    product_rows = [row for row in csv_rows if row["kind"] == "product"]
    green = sum(similarity_color(row) == "catalog-high" for row in product_rows)
    remaining = [row for row in csv_rows if needs_review(row)]
    summary = dict(catalog_unique_names=len(entries), catalog_files=len(workbooks), workbooks=workbooks,
                   catalog_conflict_count=len(conflicts), catalog_conflicts=conflicts, photos=len(photos),
                   after_product_statuses=dict(totals), products=sum(p["after_stats"]["product_rows"] for p in photos),
                   pass_similarity=PASS_SIMILARITY, similarity_passing=green,
                   similarity_below_threshold=len(product_rows)-green,
                   remaining_review_count=len(remaining), remaining_rows=remaining,
                   comparisons=[{key: p[key] for key in ("file", "before_stats", "after_stats", "segmentation")} for p in photos])
    (ROOT / "eval_outputs/ocr_catalog_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    with (ROOT / "eval_outputs/ocr_catalog_matches.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(csv_rows[0])); writer.writeheader(); writer.writerows(csv_rows)
    with (ROOT / "eval_outputs/ocr_catalog_remaining.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(csv_rows[0])); writer.writeheader(); writer.writerows(remaining)
    write_excel(csv_rows, ROOT / "eval_outputs/ocr_catalog_matches.xlsx")
    (ROOT / "eval_outputs/ocr_catalog_report.html").write_text(
        '<!doctype html><meta charset="utf-8"><title>OCR и справочники</title><style>body{font:15px system-ui;margin:24px;background:#f5f5f5}summary{font-weight:bold;font-size:18px;cursor:pointer}'
        'details{padding:15px;background:white;margin:16px 0}table{border-collapse:collapse;min-width:1100px}td,th{border:1px solid #ccc;padding:7px;text-align:left;vertical-align:top}'
        'th{background:#e4e8ed}.scroll{overflow:auto}.catalog-low{background:#fff2d3}'
        '.catalog-high{background:#e4f5e4}td:nth-child(2),td:nth-child(6){min-width:230px}'
        '.controls{position:sticky;top:0;background:#f5f5f5;padding:12px 0;border-bottom:1px solid #bbb;z-index:1}'
        'select,input{font:inherit;padding:8px;margin:4px}#visible-count{margin-left:12px}</style>'
        '<h1>Распознавание и сопоставление со справочниками</h1><p>RapidOCR / Cyrillic PP-OCRv5. '
        f'{len(photos)} фото; {len(entries)} уникальных названий в {len(workbooks)} Excel-файлах; '
        f'конфликтов артикулов: {len(conflicts)}.</p>'
        f'<p><strong>{summary["products"]} товарных строк: {totals["exact_name"]} совпадений названий после нормализации, '
        f'{totals["suggested"]} вероятных вариантов, {totals["review"]} требуют проверки.</strong></p>'
        f'<p><strong>По цвету: {green} строк со сходством {PASS_SIMILARITY:g}% и выше — зелёные; {len(product_rows)-green} строк со сходством ниже {PASS_SIMILARITY:g}% — жёлтые. '
        f'Всего для проверки: {len(remaining)} строк.</strong> Это список возможных ошибок и неоднозначностей, а не число подтверждённых ошибок.</p>'
        '<p>Сходство со справочником — не процент правильности OCR. exact_name — точное нормализованное название; suggested — вероятный кандидат; '
        'review — проверить; not_found — убедительный кандидат не найден. НДС, доставка и итог исключены из товарного сопоставления. '
        'Количество и цена сохранены как прочитаны: в справочнике нет цен, а складской остаток не равен количеству в чеке. '
        'Привязка исходного OCR к строкам приблизительная; нужна ручная проверка спорных и объединённых строк.</p>'
        '<p>Нормализация учитывает регистр, пробелы, знаки размеров и похожие кириллические/латинские буквы. '
        'Совпадение названия не подтверждает цену, количество или единственность артикула.</p>'
        '<p>Уверенность OCR имени — среднее confidence только текстовых блоков названия, умноженное на 100. '
        'Это оценка движка, а не проверенная вероятность правильного распознавания и не сходство со справочником. '
        'Зелёная строка может оставаться спорной из-за чисел, близких вариантов или нескольких артикулов: причина указана отдельно.</p>'
        '<p><a href="ocr_catalog_matches.xlsx">Открыть Excel (есть лист «Проверить»)</a> · <a href="ocr_catalog_matches.csv">Все строки CSV</a> · '
        '<a href="ocr_catalog_remaining.csv">Оставшиеся спорные строки CSV</a> · <a href="ocr_catalog_summary.json">Сводка JSON</a></p>'
        '<div class="controls"><label>Показать <select id="row-filter"><option value="products">Все товары</option>'
        f'<option value="low">Сходство ниже {PASS_SIMILARITY:g}% ({len(product_rows)-green})</option><option value="pending">Все спорные ({len(remaining)})</option>'
        '<option value="all">Все строки, включая НДС и доставку</option></select></label>'
        '<input id="search" type="search" placeholder="Поиск по товару или фото" aria-label="Поиск по товару или фото"><span id="visible-count"></span></div>'
        + ''.join(sections) + '''<script>
function filterRows(){
 const mode=document.getElementById('row-filter').value;
 const search=document.getElementById('search').value.toLocaleLowerCase(); let total=0;
 document.querySelectorAll('details.photo').forEach(section=>{
  const photo=section.querySelector('summary').textContent.toLocaleLowerCase(); let count=0;
  section.querySelectorAll('tbody tr').forEach(row=>{
   const product=row.dataset.kind==='product';
   const selected=mode==='all'||(mode==='products'&&product)||(mode==='low'&&product&&row.classList.contains('catalog-low'))||(mode==='pending'&&row.dataset.pending==='true');
   const visible=selected&&(!search||(photo+' '+row.textContent.toLocaleLowerCase()).includes(search));
   row.hidden=!visible; if(visible) count++;
  });
  section.hidden=count===0; total+=count;
 });
 document.getElementById('visible-count').textContent='Показано строк: '+total;
}
document.getElementById('row-filter').addEventListener('change',filterRows);
document.getElementById('search').addEventListener('input',filterRows); filterRows();
</script>''', encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False)[:1800])


if __name__ == "__main__":
    main()
